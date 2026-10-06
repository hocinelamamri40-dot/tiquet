import json, sys
from pathlib import Path
from datetime import datetime, date
import openpyxl

ROOT = Path(__file__).parent
def value(v):
    return v.isoformat() if isinstance(v, (datetime, date)) else v
def run(source):
    workbook = openpyxl.load_workbook(source, data_only=True)
    folder = ROOT / 'data'
    folder.mkdir(exist_ok=True)
    archive = {s.title: [[value(v) for v in row] for row in s.iter_rows(values_only=True)] for s in workbook}
    (folder / 'excel_archive.json').write_text(json.dumps(archive, ensure_ascii=False, indent=2), encoding='utf-8')
    tickets = []
    for row_number, row in enumerate(workbook['SUIVI CURATIF'].iter_rows(min_row=2, values_only=True), 2):
        if not any(row[:6]):
            continue
        equipment, kind, reason, intervention, comment, status = row[:6]
        raw_status = str(status or '').strip()
        normalized = 'Clos' if raw_status.lower() in ('clos', 'clôturé', 'cloture', 'fermé') else raw_status or 'Ouvert'
        tickets.append(dict(id=f'TKT-{len(tickets)+1:04}', equipment=str(equipment or ''), type=str(kind or ''), title=str(reason or ''), date=value(intervention)[:10] if isinstance(intervention, (date, datetime)) else '', originalDate=value(intervention), comment=str(comment or ''), status=normalized, sourceRow=row_number, originalStatus=raw_status))
    destination = folder / 'tickets.json'
    if destination.exists():
        raise RuntimeError('tickets.json existe déjà : import annulé pour préserver vos modifications.')
    destination.write_text(json.dumps(tickets, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'imported':len(tickets), 'statuses':{s:sum(t['status']==s for t in tickets) for s in sorted(set(t['status'] for t in tickets))}}, ensure_ascii=True))
if __name__ == '__main__':
    run(sys.argv[1])
